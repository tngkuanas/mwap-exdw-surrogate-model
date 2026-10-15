"""
Submission Readiness and End-to-End Delivery Verification Tests.
ExxonMobil DataWorks Challenge 2026.
"""
import json
from pathlib import Path
import numpy as np
import openpyxl
import pytest

from ensemble.submission_pipeline import (
    validate_deliverable_workbook,
    audit_rate_reconciliation,
    RS_SOLUTION,
    EXPECTED_COLUMNS,
)


def _resolve_test_file(rel_path: str) -> Path:
    """Resolves test artifacts dynamically regardless of current working directory."""
    candidates = [
        Path(rel_path),
        Path("Anas-exp") / rel_path,
        Path(__file__).resolve().parent.parent / rel_path,
        Path(__file__).resolve().parent.parent.parent / rel_path,
    ]
    for c in candidates:
        if c.exists():
            return c
    return Path(rel_path)


def test_locked_config_integrity():
    config_path = _resolve_test_file("src/ensemble/locked_ensemble_config.json")
    assert config_path.exists(), f"locked_ensemble_config.json does not exist at {config_path}!"
    
    with open(config_path) as f:
        config = json.load(f)
        
    assert config["status"] == "FROZEN"
    assert config["model_name"] == "Blend_Phase_Specific_Simplex"
    
    w_oil = config["weights"]["oil"]
    w_wat = config["weights"]["water"]
    
    assert abs(sum(w_oil.values()) - 1.0) < 1e-6
    assert abs(sum(w_wat.values()) - 1.0) < 1e-6
    assert all(v >= 0.0 for v in w_oil.values())
    assert all(v >= 0.0 for v in w_wat.values())
    
    # Specific frozen weights
    assert abs(w_oil["Recursive_ExtraTrees"] - 0.240466) < 1e-4
    assert abs(w_oil["Learned_Exponential_DCA"] - 0.617902) < 1e-4
    assert abs(w_oil["Learned_Hyperbolic_DCA"] - 0.0) < 1e-6
    assert abs(w_oil["Hybrid_Dynamic"] - 0.141631) < 1e-4
    
    assert abs(w_wat["Recursive_ExtraTrees"] - 0.497170) < 1e-4
    assert abs(w_wat["Learned_Exponential_DCA"] - 0.251415) < 1e-4
    assert abs(w_wat["Learned_Hyperbolic_DCA"] - 0.251415) < 1e-4
    
    # Physics constants
    assert config["physics"]["solution_gas_oil_ratio_rs_mscf_per_stb"] == RS_SOLUTION
    
    # Horizons
    assert config["forecast_horizons"]["Case 1"]["num_quarters"] == 40
    for i in range(2, 11):
        assert config["forecast_horizons"][f"Case {i}"]["num_quarters"] == 80


def test_official_deliverable_workbook_verification():
    file_path = _resolve_test_file("outputs/submission/09_Template_Deliverable_Submission.xlsx")
    assert file_path.exists(), f"Official deliverable file missing at {file_path}!"
    
    audit = validate_deliverable_workbook(file_path)
    assert audit["all_checks_passed"] is True
    assert audit["total_sheets"] == 10
    assert audit["nan_count"] == 0
    assert audit["negative_values_count"] == 0
    assert audit["monotonicity_violations"] == 0
    assert audit["gor_discrepancies"] == 0
    
    # Check Case 1 horizon
    assert audit["sheets"]["Case 1"]["num_quarters"] == 40
    assert audit["sheets"]["Case 1"]["oil_monotonic"] is True
    assert audit["sheets"]["Case 1"]["water_monotonic"] is True
    assert audit["sheets"]["Case 1"]["gas_monotonic"] is True
    
    # Check Cases 2-10 horizons
    for i in range(2, 11):
        sheet_key = f"Case {i}"
        assert audit["sheets"][sheet_key]["num_quarters"] == 80
        assert audit["sheets"][sheet_key]["oil_monotonic"] is True
        assert audit["sheets"][sheet_key]["water_monotonic"] is True
        assert audit["sheets"][sheet_key]["gas_monotonic"] is True
        assert audit["sheets"][sheet_key]["max_gor_residual"] < 1e-4


def test_smoke_test_deliverable_verification():
    file_path = _resolve_test_file("outputs/submission/smoke_test_dev_cases_1_to_10.xlsx")
    assert file_path.exists(), f"Smoke test deliverable file missing at {file_path}!"
    
    audit = validate_deliverable_workbook(file_path)
    assert audit["all_checks_passed"] is True
    assert audit["total_sheets"] == 10
    assert audit["nan_count"] == 0
    assert audit["negative_values_count"] == 0
    assert audit["monotonicity_violations"] == 0


def test_rate_reconciliation_exactness():
    rates = np.array([5000.0, 4800.0, 4600.0, 4400.0, 4200.0])
    anchor_cum = 38000000.0
    dt_days = 91.3125
    cums = anchor_cum + np.cumsum(rates * dt_days)
    
    err = audit_rate_reconciliation(rates, cums, anchor_cum, dt_days)
    assert err < 1e-10, f"Rate reconciliation error exceeds tolerance: {err}"


def test_submission_audit_report_contents():
    report_path = _resolve_test_file("outputs/submission/submission_audit_report.json")
    assert report_path.exists(), f"submission_audit_report.json missing at {report_path}!"
    
    with open(report_path) as f:
        report = json.load(f)
        
    assert report["status"] == "VERIFIED_SUBMISSION_READY"
    assert 0.0 < report["rate_reconciliation_max_error"] < 1e-6
    assert 0.0 < report["solution_gor_residual_max"] < 1e-4
    assert np.isfinite(report["rate_reconciliation_max_error"])
    assert np.isfinite(report["solution_gor_residual_max"])
    assert report["tail_stabilization_audit"]["case_2_q80_cum_stb"] > report["tail_stabilization_audit"]["case_2_q40_cum_stb"]
    assert "scenarios" in report
    assert len(report["scenarios"]) == 10
    for sc_info in report["scenarios"]:
        assert sc_info["decline_sensitivity"] == 1.0


def test_scenario_loader_contract():
    from ensemble.scenario_loader import (
        load_production_scenarios,
        ProductionScenarioConfig,
        PARAMETER_BOUNDS,
    )
    import pandas as pd
    
    scenarios = load_production_scenarios()
    assert len(scenarios) == 10
    
    expected_cands = [56, 440, 242, 461, 95, 286, 392, 119, 290, 144]
    for idx, sc in enumerate(scenarios):
        assert sc.case_id == f"Case {idx + 1}"
        assert sc.candidate_id == expected_cands[idx]
        assert sc.decline_sensitivity == 1.0
        if idx == 0:
            assert sc.horizon_quarters == 40
        else:
            assert sc.horizon_quarters == 80
            
        for p_name, (low, high) in PARAMETER_BOUNDS.items():
            val = float(sc.params[p_name])
            assert low <= val <= high, f"{p_name}={val} out of [{low}, {high}]"
            
    # Test strict contract enforcement
    valid_params = pd.Series({
        "Fault Transmissibility": 0.1,
        "Porosity Multiplier": 1.0,
        "Permeability Multiplier": 2.0,
        "Aquifer Pore Volume": 100.0,
    })
    
    # Must reject sensitivity != 1.0
    with pytest.raises(ValueError, match="decline_sensitivity = 1.0"):
        ProductionScenarioConfig(
            case_id="Case 1",
            candidate_id=999,
            params=valid_params,
            horizon_quarters=40,
            selection_strategy="Test",
            target_role="Test",
            history_match_metrics={},
            research_source_artifact="test",
            provenance_commit="test",
            decline_sensitivity=0.85,
        )
        
    # Must reject out-of-bounds parameters
    bad_params = valid_params.copy()
    bad_params["Fault Transmissibility"] = 2.5  # bound is [0.01, 1.0]
    with pytest.raises(ValueError, match="out of authoritative bounds"):
        ProductionScenarioConfig(
            case_id="Case 1",
            candidate_id=999,
            params=bad_params,
            horizon_quarters=40,
            selection_strategy="Test",
            target_role="Test",
            history_match_metrics={},
            research_source_artifact="test",
            provenance_commit="test",
            decline_sensitivity=1.0,
        )


def test_submission_deliverable_consistency_with_research():
    deliverable_path = _resolve_test_file("outputs/submission/09_Template_Deliverable_Submission.xlsx")
    wb = openpyxl.load_workbook(deliverable_path, data_only=True)
    
    # Case 1 should have 40 rows of forecast
    ws1 = wb["Case 1"]
    row_count_1 = sum(1 for row in ws1.iter_rows(min_row=2, max_col=1, values_only=True) if row[0] is not None)
    assert row_count_1 == 40
    # Terminal oil should be ~47.65 MMSTB (un-skewed physics, not 45.989 which was with 0.85 boost or cut)
    final_oil_1 = ws1.cell(row=41, column=3).value
    assert abs(final_oil_1 - 47650698.0) < 1000.0, f"Expected ~47.65 MMSTB, got {final_oil_1}"
    
    # Case 2 should have 80 rows of forecast
    ws2 = wb["Case 2"]
    row_count_2 = sum(1 for row in ws2.iter_rows(min_row=2, max_col=1, values_only=True) if row[0] is not None)
    assert row_count_2 == 80
    final_oil_2 = ws2.cell(row=81, column=3).value
    assert abs(final_oil_2 - 51172196.0) < 1000.0, f"Expected ~51.17 MMSTB, got {final_oil_2}"


def test_deliverable_exact_forecast_dates_dimensions_and_finite_cells():
    """
    Validates exact sheet dimensions, column headers, forecast date sequences matching template,
    and strict finiteness and non-negativity across all 2,280 numeric forecast cells.
    """
    from ensemble.submission_pipeline import resolve_template_path
    tmpl_path = resolve_template_path()
    deliv_path = _resolve_test_file("outputs/submission/09_Template_Deliverable_Submission.xlsx")
    assert tmpl_path.exists(), f"Official template missing at {tmpl_path}"
    assert deliv_path.exists(), f"Official deliverable missing at {deliv_path}"
    
    wb_tmpl = openpyxl.load_workbook(tmpl_path, data_only=True)
    wb_deliv = openpyxl.load_workbook(deliv_path, data_only=True)
    
    assert set(wb_deliv.sheetnames) == {f"Case {i}" for i in range(1, 11)}
    total_finite_numeric_cells = 0
    
    for case_num in range(1, 11):
        sheet_name = f"Case {case_num}"
        ws_tmpl = wb_tmpl[sheet_name]
        ws_deliv = wb_deliv[sheet_name]
        
        rows_tmpl = list(ws_tmpl.iter_rows(values_only=True))
        rows_deliv = list(ws_deliv.iter_rows(values_only=True))
        
        expected_rows = 41 if case_num == 1 else 81
        assert len(rows_deliv) == expected_rows, f"{sheet_name} dimension error: {len(rows_deliv)} != {expected_rows}"
        assert len(rows_tmpl) == expected_rows, f"{sheet_name} template dimension error"
        
        # Check header
        assert list(rows_deliv[0][:4]) == EXPECTED_COLUMNS
        
        # Check row by row: date and 3 numeric cells
        for r_idx, (rt, rd) in enumerate(zip(rows_tmpl[1:], rows_deliv[1:]), start=2):
            # Exact forecast date match
            date_t = str(rt[0])
            date_d = str(rd[0])
            assert date_t == date_d, f"{sheet_name} row {r_idx} date mismatch: {date_d} != {date_t}"
            
            # Validate 3 fluid cumulatives: gas, oil, water
            for col_idx, (c_name, val) in enumerate(zip(["gas", "oil", "water"], rd[1:4]), start=2):
                assert val is not None, f"{sheet_name} row {r_idx} {c_name} is None"
                assert isinstance(val, (int, float)), f"{sheet_name} row {r_idx} {c_name} is not numeric: {type(val)}"
                assert np.isfinite(val), f"{sheet_name} row {r_idx} {c_name} is non-finite: {val}"
                assert not np.isnan(val), f"{sheet_name} row {r_idx} {c_name} is NaN"
                assert val >= 0.0, f"{sheet_name} row {r_idx} {c_name} is negative: {val}"
                total_finite_numeric_cells += 1
                
    assert total_finite_numeric_cells == 2280, f"Expected 2280 finite cells, counted {total_finite_numeric_cells}"


def test_archive_case_isolation():
    """
    Verifies that Cases 86-100 blind test case production records are strictly excluded from
    packaged archive data. Rejects any record with simulation reservoir case ID in [86, 100].
    Explicitly distinguishes simulation reservoir case IDs (1-100) from scenario candidate
    identifiers (e.g., Cand 440, 242, 461) which represent Iteration 4 MCMC parameter draws.
    """
    import io
    import zipfile
    import pandas as pd
    
    # Check 1: If archive zip exists, inspect contents directly inside ZIP
    zip_candidates = [
        Path("outputs/submission/MWAP_Energy_Analytics.zip"),
        Path("outputs/MWAP_Energy_Analytics.zip"),
        _resolve_test_file("outputs/submission/MWAP_Energy_Analytics.zip"),
    ]
    zip_path = next((p for p in zip_candidates if p.exists()), None)
    
    # Check 2: If running inside extracted package without ZIP, inspect extracted parquets
    local_sim = _resolve_test_file("data/simulation_timeseries.parquet")
    local_prod = _resolve_test_file("data/shared/production_timeseries.parquet")
    
    tested = False
    
    if zip_path is not None:
        with zipfile.ZipFile(zip_path, "r") as zf:
            names = zf.namelist()
            sim_names = [n for n in names if n.endswith("data/simulation_timeseries.parquet")]
            prod_names = [n for n in names if n.endswith("data/shared/production_timeseries.parquet")]
            
            if sim_names:
                df_sim_zip = pd.read_parquet(io.BytesIO(zf.read(sim_names[0])))
                sim_cases = df_sim_zip["case_id"].unique()
                leaked_sim = set(sim_cases).intersection(set(range(86, 101)))
                assert len(leaked_sim) == 0, f"BLIND CASE LEAK in archive ZIP simulation_timeseries: cases {leaked_sim}"
                assert int(df_sim_zip["case_id"].max()) <= 85
                tested = True
                
            if prod_names:
                df_prod_zip = pd.read_parquet(io.BytesIO(zf.read(prod_names[0])))
                prod_cases = df_prod_zip["case_num"].unique()
                leaked_prod = set(prod_cases).intersection(set(range(86, 101)))
                assert len(leaked_prod) == 0, f"BLIND CASE LEAK in archive ZIP production_timeseries: cases {leaked_prod}"
                assert int(df_prod_zip["case_num"].max()) <= 85
                tested = True
                
    if local_sim.exists() and local_prod.exists():
        df_sim = pd.read_parquet(local_sim)
        # If running in filtered environment (max <= 85), assert zero blind cases
        if int(df_sim["case_id"].max()) <= 85:
            sim_cases = df_sim["case_id"].unique()
            leaked_sim = set(sim_cases).intersection(set(range(86, 101)))
            assert len(leaked_sim) == 0, f"Blind cases in local simulation_timeseries: {leaked_sim}"
            
            df_prod = pd.read_parquet(local_prod)
            prod_cases = df_prod["case_num"].unique()
            leaked_prod = set(prod_cases).intersection(set(range(86, 101)))
            assert len(leaked_prod) == 0, f"Blind cases in local production_timeseries: {leaked_prod}"
            assert int(df_prod["case_num"].max()) <= 85
            tested = True
            
    if not tested:
        pytest.skip("Archive ZIP not yet generated; run 'python run_submission.py --build-archive' to test.")
        
    # Disambiguate scenario candidate IDs vs reservoir case IDs
    report_path = _resolve_test_file("outputs/submission/submission_audit_report.json")
    if report_path.exists():
        with open(report_path) as f:
            rep = json.load(f)
        candidate_ids = [sc["candidate_id"] for sc in rep["scenarios"]]
        assert 440 in candidate_ids, "Candidate 440 must be recognized as valid MCMC scenario parameter draw"
        assert any(cid > 100 for cid in candidate_ids), "Candidate IDs > 100 must not be confused with reservoir case IDs"

